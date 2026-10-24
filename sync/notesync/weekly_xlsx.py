"""週報の Excel(F-29)。表ごとに1シート。管理者が並べ替えたり絞り込んだりできるよう、見出しにフィルタをつける。

openpyxl がない PC では読み込まれず、Markdown だけを作る(weekly.py)。
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def write_xlsx(path: Path, report) -> None:
    wb = Workbook()
    first = wb.active
    first.title = "見込み"
    for line in report.lines:
        if line.startswith("- "):
            k, _, v = line[2:].partition(": ")
            first.append([k.strip("*"), v.strip("*")])
        elif line.startswith("# "):
            first.append([line[2:]])
            first.cell(row=first.max_row, column=1).font = Font(bold=True, size=14)
        if line.startswith("## 出荷準備度"):
            break
    first.column_dimensions["A"].width = 28
    first.column_dimensions["B"].width = 60
    head = PatternFill("solid", fgColor="DDEBF7")
    for name, rows in report.rows.items():
        ws = wb.create_sheet(name[:31])
        for r in rows:
            ws.append(r)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = head
        if len(rows) > 1:
            ws.auto_filter.ref = ws.dimensions
        else:
            ws.append(["なし"])
        ws.freeze_panes = "A2"
        for i, col in enumerate(zip(*rows), start=1):
            width = max(len(str(v)) for v in col) * 2 + 2
            ws.column_dimensions[get_column_letter(i)].width = min(max(width, 10), 60)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
